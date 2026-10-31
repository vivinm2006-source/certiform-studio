import csv
import io
import re
import datetime
from typing import List, Dict, Any, Tuple
import openpyxl
from sqlalchemy.orm import Session
from app.models import Recipient, FormField, Project

def clean_corrupted_s(text: str) -> str:
    """
    Clean corrupted letter 's' substitutions where Cyrillic 'з'/'З' or digit '3'
    was entered in place of 's'/'S' (due to broken keyboard, OCR, or leetspeak).
    Carefully preserves valid numbers like dates, phone numbers, ordinals, and years.
    """
    if not text or not isinstance(text, str):
        return text

    res = text
    # 0. Check if this contains an email address
    is_email = '@' in res

    # 1. Protect valid ordinals like 3rd, 3RD, 3th, 3nd, 3st
    res = re.sub(r'\b3(rd|RD|th|TH|nd|ND|st|ST)\b', r'__ORD_3\1__', res)

    # 2. Cyrillic Ze (U+0437 lowercase and U+0417 uppercase)
    # If at start of a capitalized word and not email, convert to 'S'
    if not is_email:
        res = re.sub(r'\b[зЗ]([A-Za-z]*)', lambda m: 'S' + m.group(1), res)
    res = res.replace('\u0437', 's').replace('\u0417', 'S')
    res = res.replace('з', 's').replace('З', 'S')

    # 3. Handle multiple 3s in uppercase/lowercase words (e.g. ADDRE33 -> ADDRESS, addre33 -> address)
    def replace_consecutive_3(m):
        before = m.group(1)
        threes = m.group(2)
        rep = 'S' if before.isupper() else 's'
        return before + (rep * len(threes))
    res = re.sub(r'([A-Za-z])(3+)', replace_consecutive_3, res)

    # 4. Handle 3 between letters (e.g. TIME3TAMP -> TIMESTAMP, IN3TITUTION -> INSTITUTION, mc3ankar -> mcsankar)
    def replace_mid_3(m):
        before = m.group(1)
        after = m.group(2)
        rep = 'S' if (before.isupper() and after.isupper()) else 's'
        return f"{before}{rep}{after}"
    prev = None
    while prev != res:
        prev = res
        res = re.sub(r'([A-Za-z])3([A-Za-z])', replace_mid_3, res)

    # 5. Handle 3 followed by letters at start of word or email (e.g. 3anthosh -> Santhosh / santhosh, 3ukanth -> Sukanth)
    def replace_start_3(m):
        word = m.group(0)
        if word.isdigit():
            return word
        if word.startswith('3'):
            first_rep = 's' if is_email else 'S'
            word = first_rep + word[1:]
        return word
    res = re.sub(r'\b3[A-Za-z]+', replace_start_3, res)

    # 6. Normalize common corrupted student names like 'SREE sukanth' -> 'SREE Sukanth'
    if 'SREE' in res and 'sukanth' in res:
        res = res.replace('SREE sukanth', 'SREE Sukanth')
    elif 'sree' in res.lower() and 'sukanth' in res.lower():
        res = re.sub(r'\bsree\s+sukanth\b', 'Sree Sukanth', res, flags=re.IGNORECASE)

    # 7. Restore protected ordinals
    res = re.sub(r'__ORD_3([A-Za-z]+)__', r'3\1', res)

    return res

    return res

def normalize_spaces(text: str) -> str:
    """
    Strip leading/trailing whitespace and collapse internal consecutive spaces
    or whitespace gaps to a single space.
    """
    if text is None:
        return ""
    if not isinstance(text, str):
        text = str(text)
    return re.sub(r'[ \t\r\n]+', ' ', text).strip()

def clean_record_dict(d: Dict[str, Any]) -> Dict[str, Any]:
    """Clean all keys and string values in a recipient record dictionary."""
    if not isinstance(d, dict):
        return d
    cleaned = {}
    for k, v in d.items():
        k_clean = normalize_spaces(clean_corrupted_s(str(k)))
        if isinstance(v, str):
            v_clean = normalize_spaces(clean_corrupted_s(v))
        else:
            v_clean = v
        cleaned[k_clean] = v_clean
    return cleaned

SYNONYM_MAP = {
    # Recipient & Student name synonyms
    "name": "recipient_name",
    "full_name": "recipient_name",
    "fullname": "recipient_name",
    "student_name": "recipient_name",
    "student": "recipient_name",
    "students_name": "recipient_name",
    "student's_name": "recipient_name",
    "s_name": "recipient_name",
    "sname": "recipient_name",
    "s.name": "recipient_name",
    "s_word": "recipient_name",
    "name_of_the_student": "recipient_name",
    "name_of_student": "recipient_name",
    "name_of_the_candidate": "recipient_name",
    "name_of_candidate": "recipient_name",
    "candidate_name": "recipient_name",
    "candidate": "recipient_name",
    "participant_name": "recipient_name",
    "participant": "recipient_name",
    "attendee_name": "recipient_name",
    "attendee": "recipient_name",
    "recipient": "recipient_name",
    "recipient_name": "recipient_name",
    "person_name": "recipient_name",
    "person": "recipient_name",
    "member_name": "recipient_name",
    "member": "recipient_name",
    "first_name": "recipient_name",
    "firstname": "recipient_name",
    "last_name": "recipient_name",
    "username": "recipient_name",
    
    # Department / Academic Branch
    "department": "department",
    "dept": "department",
    "branch": "department",
    "stream": "department",
    "discipline": "department",
    "major": "department",
    "department_name": "department",
    "dept_name": "department",
    "degree_branch": "department",

    # Date of Birth (DOB) synonyms
    "dob": "dob",
    "d_o_b": "dob",
    "d.o.b": "dob",
    "d.o.b.": "dob",
    "date_of_birth": "dob",
    "birth_date": "dob",
    "birthdate": "dob",
    "birthday": "dob",
    "date of birth": "dob",
    "birth date": "dob",

    # Email synonyms
    "mail": "email",
    "email_address": "email",
    "e_mail": "email",
    "emailaddress": "email",
    
    # Phone synonyms
    "phone": "phone_number",
    "mobile": "phone_number",
    "contact": "phone_number",
    
    # Registration & Credential ID
    "registration_no": "registration_id",
    "reg_no": "registration_id",
    "reg_id": "registration_id",
    "roll_no": "registration_id",
    "rollno": "registration_id",
    "id": "registration_id",
    "certificate_id": "certificate_id",

    # Organization / College / Institution
    "organization": "institution",
    "org": "institution",
    "company": "institution",
    "college": "institution",
    "university": "institution",
    "school": "institution",
    "institute": "institution",
    "institution": "institution",
    "venue": "venue",
    "location": "venue",
    
    # Award / Course / Event / Session Topic
    "award": "award_category",
    "category": "award_category",
    "course": "course_name",
    "course_name": "course_name",
    "date": "event_date",
    "event_date": "event_date",
    "event": "event_name",
    "event_name": "event_name",
    "session": "event_name",
    "session_topic": "event_name",
    "topic": "event_name",
    "theme": "event_name",
    "workshop": "event_name",
    "conference": "event_name",
}

def suggest_column_mapping(headers: List[str], available_keys: List[str]) -> Dict[str, str]:
    """
    Intelligently map spreadsheet headers to form field binding keys.
    Uses exact, lowercase, normalized, and synonym matching.
    """
    mapping = {}
    normalized_keys = {k.lower().replace(" ", "_").replace("-", "_"): k for k in available_keys}
    
    for h in headers:
        if not h:
            continue
        h_clean = h.strip()
        h_norm = h_clean.lower().replace(" ", "_").replace("-", "_")
        
        # 1. Exact match with available keys
        if h_clean in available_keys:
            mapping[h_clean] = h_clean
        # 2. Normalized match
        elif h_norm in normalized_keys:
            mapping[h_clean] = normalized_keys[h_norm]
        # 3. Synonym dictionary lookup
        elif h_norm in SYNONYM_MAP and SYNONYM_MAP[h_norm] in available_keys:
            mapping[h_clean] = SYNONYM_MAP[h_norm]
        else:
            # 4. Partial substring match
            matched = False
            for norm_k, orig_k in normalized_keys.items():
                if norm_k in h_norm or h_norm in norm_k:
                    mapping[h_clean] = orig_k
                    matched = True
                    break
            if not matched:
                mapping[h_clean] = ""  # Let user choose or ignore
                
    return mapping

def parse_csv_content(content_bytes: bytes) -> Tuple[List[str], List[Dict[str, Any]]]:
    """Parse CSV bytes and return headers and raw row dictionaries."""
    try:
        text = content_bytes.decode("utf-8-sig")
    except UnicodeDecodeError:
        text = content_bytes.decode("latin-1")
        
    reader = csv.reader(io.StringIO(text))
    rows = list(reader)
    if not rows:
        return [], []
        
    raw_headers = rows[0]
    last_col = -1
    for i, h in enumerate(raw_headers):
        if h and h.strip():
            last_col = i
    if last_col == -1:
        return [], []
        
    headers = [normalize_spaces(clean_corrupted_s(raw_headers[i])) if i < len(raw_headers) and raw_headers[i].strip() else f"Column_{i+1}" for i in range(last_col + 1)]
    data_rows = []
    
    for row_vals in rows[1:]:
        if not any(v.strip() for v in row_vals):
            continue  # skip empty lines
        row_dict = {}
        for idx, h in enumerate(headers):
            val = row_vals[idx].strip() if idx < len(row_vals) else ""
            row_dict[h] = normalize_spaces(clean_corrupted_s(val))
        data_rows.append(row_dict)
        
    return headers, data_rows

def parse_xlsx_content(content_bytes: bytes) -> Tuple[List[str], List[Dict[str, Any]]]:
    """Parse XLSX bytes and return headers and raw row dictionaries."""
    wb = openpyxl.load_workbook(filename=io.BytesIO(content_bytes), data_only=True)
    sheet = wb.active
    
    headers = []
    data_rows = []
    
    for row_idx, row in enumerate(sheet.iter_rows(values_only=True)):
        if row_idx == 0:
            last_col = -1
            for i, cell in enumerate(row):
                if cell is not None and str(cell).strip():
                    last_col = i
            if last_col == -1:
                return [], []
            headers = [normalize_spaces(clean_corrupted_s(str(row[i]))) if i < len(row) and row[i] is not None and str(row[i]).strip() else f"Column_{i+1}" for i in range(last_col + 1)]
        else:
            if not any(cell is not None and str(cell).strip() != "" for cell in row):
                continue
            row_dict = {}
            for col_idx, h in enumerate(headers):
                cell_val = row[col_idx] if col_idx < len(row) else ""
                if cell_val is None:
                    row_dict[h] = ""
                elif isinstance(cell_val, (datetime.datetime, datetime.date)):
                    row_dict[h] = cell_val.strftime("%Y-%m-%d")
                elif isinstance(cell_val, float) and cell_val.is_integer():
                    row_dict[h] = str(int(cell_val))
                else:
                    row_dict[h] = normalize_spaces(clean_corrupted_s(str(cell_val)))
            data_rows.append(row_dict)
            
    return headers, data_rows

def generate_empty_excel_template(form_fields: List[FormField]) -> bytes:
    """Generate empty Excel template matching the project's active form fields."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Recipients"
    
    headers = [field.label for field in form_fields]
    if not headers:
        headers = ["Recipient Name", "Email", "Institution", "Award Category"]
        
    ws.append(headers)
    
    # Add example row
    sample_row = []
    for f in form_fields:
        if "name" in f.binding_key:
            sample_row.append("Ananya Kumar")
        elif "email" in f.binding_key:
            sample_row.append("ananya@example.com")
        elif "date" in f.binding_key:
            sample_row.append("2026-10-27")
        elif "id" in f.binding_key:
            sample_row.append("REG-2026-001")
        else:
            sample_row.append(f.default_value or "Sample Value")
    if sample_row:
        ws.append(sample_row)
        
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
